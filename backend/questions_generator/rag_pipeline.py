import os
import json
import re

from dotenv import load_dotenv

from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_groq import ChatGroq


# ============================================================
# ENVIRONMENT
# ============================================================

load_dotenv(
    os.path.join(
        os.path.dirname(os.path.dirname(__file__)),
        ".env"
    )
)


# ============================================================
# LLM
# ============================================================

def create_llm():

    return ChatGroq(
        model="openai/gpt-oss-120b",
        temperature=0.3
    )


# ============================================================
# LOAD RESUME
# ============================================================

def load_resume(resume_path):
    if resume_path.lower().endswith(".pdf"):
        return PyPDFLoader(resume_path).load()

    if resume_path.lower().endswith(".docx"):
        from docx import Document
        from langchain_core.documents import Document as LangchainDocument

        document = Document(resume_path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        return [LangchainDocument(page_content=text)]

    raise ValueError("Question generation supports PDF and DOCX resumes")


# ============================================================
# CREATE VECTOR STORE
# ============================================================

def create_vector_store(documents):

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=100
    )

    chunks = splitter.split_documents(
        documents
    )

    embeddings = HuggingFaceEmbeddings(
        model_name="sentence-transformers/all-MiniLM-L6-v2"
    )

    vector_store = FAISS.from_documents(
        chunks,
        embeddings
    )

    return vector_store


# ============================================================
# RETRIEVE RESUME CONTEXT
# ============================================================

def retrieve_resume_context(
    vector_store,
    job_description,
    k=8
):

    documents = vector_store.similarity_search(
        job_description,
        k=k
    )

    context = "\n\n".join(
        doc.page_content
        for doc in documents
    )

    return context


# ============================================================
# DIFFICULTY MARKS
# ============================================================

def get_marks(difficulty):

    if difficulty == "Easy":
        return 5

    if difficulty == "Medium":
        return 10

    if difficulty == "Hard":
        return 15

    return 5


# ============================================================
# GENERATE INTERVIEW
# ============================================================

def generate_interview(
    resume_path,
    job_description,
    mode,
    total_questions,
    easy_questions,
    medium_questions,
    hard_questions
):

    documents = load_resume(
        resume_path
    )

    vector_store = create_vector_store(
        documents
    )

    resume_context = retrieve_resume_context(
        vector_store,
        job_description
    )

    llm = create_llm()

    # --------------------------------------------------------
    # Determine question types
    # --------------------------------------------------------

    if mode == "1":
        interview_type = "MCQ"

    elif mode == "2":
        interview_type = "DESCRIPTIVE"

    else:
        interview_type = "MIXED"

    # --------------------------------------------------------
    # Prompt
    # --------------------------------------------------------

    prompt = f"""
You are an expert AI technical interviewer.

Your task is to create a personalized interview assessment
using the candidate's resume and the provided job description.

============================================================
INTERVIEW CONFIGURATION
============================================================

Interview Type:
{interview_type}

Total Questions:
{total_questions}

Easy Questions:
{easy_questions}

Medium Questions:
{medium_questions}

Hard Questions:
{hard_questions}

============================================================
QUESTION TYPE RULES
============================================================

If interview type is MCQ:

Generate ONLY MCQ questions.

Every MCQ must contain:

- question
- four options: A, B, C, D
- exactly one correct answer
- explanation
- difficulty
- marks
- type = "MCQ"

If interview type is DESCRIPTIVE:

Generate ONLY descriptive questions.

Every descriptive question must contain:

- question
- expected answer
- evaluation rubric
- difficulty
- marks
- type = "DESCRIPTIVE"

If interview type is MIXED:

Generate both MCQ and descriptive questions.

Use approximately 50% MCQ and 50% descriptive questions.

============================================================
DIFFICULTY RULES
============================================================

Easy:

Basic concepts, definitions and fundamentals.

Medium:

Application-based questions, implementation,
project understanding and moderate reasoning.

Hard:

Deep technical concepts, optimization,
architecture, debugging, trade-offs and advanced reasoning.

============================================================
IMPORTANT
============================================================

Generate EXACTLY {total_questions} questions.

Generate exactly:

Easy = {easy_questions}
Medium = {medium_questions}
Hard = {hard_questions}

Do not generate additional questions.

Questions must be relevant to the job description.

Questions should be personalized using the candidate's resume.

Do NOT invent technologies, projects, skills or experience
that are not present in the resume.

For resume-based questions, only use information actually
present in the resume.

============================================================
MARKS
============================================================

Easy = 5 marks

Medium = 10 marks

Hard = 15 marks

============================================================
MCQ RULES
============================================================

For every MCQ:

- Exactly four options.
- Exactly one correct answer.
- Avoid ambiguous options.
- Distractors should be plausible.
- Do not make the correct answer obvious because of length.
- Include an explanation for the correct answer.

============================================================
DESCRIPTIVE RUBRIC
============================================================

For descriptive questions, create a scoring rubric.

The rubric should identify the important concepts that
the candidate should mention.

The rubric will later be used by another AI evaluator.

============================================================
OUTPUT FORMAT
============================================================

Return ONLY valid JSON.

Do NOT include markdown.

Use this structure:

{{
    "questions": [
        {{
            "id": 1,
            "type": "MCQ",
            "difficulty": "Easy",
            "marks": 5,
            "question": "Question text",
            "options": {{
                "A": "Option A",
                "B": "Option B",
                "C": "Option C",
                "D": "Option D"
            }},
            "correct_answer": "B",
            "explanation": "Explanation"
        }},

        {{
            "id": 2,
            "type": "DESCRIPTIVE",
            "difficulty": "Medium",
            "marks": 10,
            "question": "Question text",
            "expected_answer": "Expected answer",
            "rubric": [
                "Important concept 1",
                "Important concept 2",
                "Important concept 3"
            ]
        }}
    ]
}}

============================================================
CANDIDATE RESUME CONTEXT
============================================================

{resume_context}

============================================================
JOB DESCRIPTION
============================================================

{job_description}
"""

    response = llm.invoke(prompt)

    content = response.content

    # --------------------------------------------------------
    # Extract JSON
    # --------------------------------------------------------

    content = clean_json_response(
        content
    )

    try:

        interview = json.loads(
            content
        )

    except json.JSONDecodeError:

        raise ValueError(
            "The LLM did not return valid JSON.\n\n"
            + content
        )

    # --------------------------------------------------------
    # Validate
    # --------------------------------------------------------

    questions = interview.get(
        "questions",
        []
    )

    if len(questions) != total_questions:

        raise ValueError(
            f"Expected {total_questions} questions "
            f"but received {len(questions)}."
        )

    validate_difficulty_distribution(
        questions,
        easy_questions,
        medium_questions,
        hard_questions
    )

    return interview


# ============================================================
# CLEAN JSON
# ============================================================

def clean_json_response(content):

    content = content.strip()

    # Remove markdown JSON fences

    content = re.sub(
        r"^```json\s*",
        "",
        content,
        flags=re.IGNORECASE
    )

    content = re.sub(
        r"^```\s*",
        "",
        content
    )

    content = re.sub(
        r"\s*```$",
        "",
        content
    )

    # Find JSON object if extra text exists

    start = content.find("{")

    end = content.rfind("}")

    if start != -1 and end != -1:

        content = content[start:end + 1]

    return content.strip()


# ============================================================
# VALIDATE DIFFICULTY
# ============================================================

def validate_difficulty_distribution(
    questions,
    easy_required,
    medium_required,
    hard_required
):

    easy_count = sum(
        1
        for q in questions
        if q["difficulty"].lower() == "easy"
    )

    medium_count = sum(
        1
        for q in questions
        if q["difficulty"].lower() == "medium"
    )

    hard_count = sum(
        1
        for q in questions
        if q["difficulty"].lower() == "hard"
    )

    if easy_count != easy_required:

        raise ValueError(
            f"Expected {easy_required} Easy questions "
            f"but received {easy_count}."
        )

    if medium_count != medium_required:

        raise ValueError(
            f"Expected {medium_required} Medium questions "
            f"but received {medium_count}."
        )

    if hard_count != hard_required:

        raise ValueError(
            f"Expected {hard_required} Hard questions "
            f"but received {hard_count}."
        )


# ============================================================
# EVALUATE DESCRIPTIVE ANSWER
# ============================================================

def evaluate_descriptive_answer(
    question,
    candidate_answer
):

    llm = create_llm()

    prompt = f"""
You are an expert technical interviewer.

Evaluate the candidate's answer against the expected answer
and rubric.

============================================================
QUESTION
============================================================

{question["question"]}

============================================================
EXPECTED ANSWER
============================================================

{question["expected_answer"]}

============================================================
RUBRIC
============================================================

{json.dumps(question["rubric"], indent=2)}

============================================================
CANDIDATE ANSWER
============================================================

{candidate_answer}

============================================================
MAXIMUM MARKS
============================================================

{question["marks"]}

============================================================
EVALUATION RULES
============================================================

Evaluate based on:

1. Technical correctness
2. Understanding of the concept
3. Completeness
4. Relevant explanation
5. Accuracy

Give partial marks when appropriate.

Do not give marks merely because the candidate uses
technical terminology.

Do not penalize different wording if the underlying
concept is correct.

Do not invent facts about the candidate.

============================================================
OUTPUT
============================================================

Return ONLY valid JSON.

{{
    "score": 0,
    "max_score": {question["marks"]},
    "feedback": "Detailed feedback",
    "strengths": [
        "Strength 1"
    ],
    "improvements": [
        "Improvement 1"
    ]
}}
"""

    response = llm.invoke(prompt)

    content = clean_json_response(
        response.content
    )

    try:

        evaluation = json.loads(
            content
        )

    except json.JSONDecodeError:

        raise ValueError(
            "Invalid JSON returned by evaluator."
        )

    # Safety validation

    score = float(
        evaluation.get(
            "score",
            0
        )
    )

    max_score = float(
        question["marks"]
    )

    score = max(
        0,
        min(
            score,
            max_score
        )
    )

    evaluation["score"] = score

    evaluation["max_score"] = max_score

    return evaluation


# ============================================================
# END
# ============================================================